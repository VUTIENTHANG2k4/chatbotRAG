from __future__ import annotations

from pathlib import Path
from typing import List, Optional

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from app.core.config import settings, VI_SEPARATORS
from app.core.logging import get_logger
from app.services.bm25_store import add_chunks_to_index, remove_document_from_index
from app.services.embeddings import embed_texts
from app.services.vectorstore import upsert_chunks, document_exists, delete_document_by_source
from app.utils.text import clean_text, extract_metadata_from_filename

logger = get_logger(__name__)

# Supported file extensions
_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".tiff", ".tif", ".bmp", ".webp"}
SUPPORTED_EXTS = {".pdf", ".docx", ".txt"} | _IMAGE_EXTS


# ── Loaders ────────────────────────────────────────────────────────────────

def load_pdf(path: Path) -> List[Document]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    docs = []

    # Try text extraction first; fall back to OCR for pages with little/no text
    _ocr_available: Optional[bool] = None

    def _ocr_available_check() -> bool:
        nonlocal _ocr_available
        if _ocr_available is not None:
            return _ocr_available
        try:
            import pytesseract  # noqa: F401
            from pdf2image import convert_from_path  # noqa: F401
            _ocr_available = True
        except ImportError:
            _ocr_available = False
            logger.warning("OCR libraries not available; install pytesseract & pdf2image")
        return _ocr_available

    def _ocr_page(page_num: int) -> str:
        """Convert one PDF page to image and run Tesseract OCR."""
        try:
            from pdf2image import convert_from_path
            import pytesseract

            images = convert_from_path(
                str(path),
                first_page=page_num,
                last_page=page_num,
                dpi=300,
            )
            if not images:
                return ""
            return pytesseract.image_to_string(images[0], lang="vie+eng")
        except Exception as e:
            logger.warning("OCR failed on page %d of %s: %s", page_num, path.name, e)
            return ""

    # Threshold: if extracted text is below this many chars, assume it's a scan
    MIN_TEXT_CHARS = 50

    for page_num, page in enumerate(reader.pages, start=1):
        text = clean_text(page.extract_text() or "")

        if len(text.strip()) < MIN_TEXT_CHARS and _ocr_available_check():
            logger.info("Page %d of %s has little text (%d chars) — attempting OCR",
                        page_num, path.name, len(text.strip()))
            ocr_text = clean_text(_ocr_page(page_num))
            if ocr_text.strip():
                text = ocr_text

        if text.strip():
            docs.append(Document(page_content=text, metadata={"page": page_num}))
        else:
            logger.debug("Page %d of %s is empty after text extraction + OCR", page_num, path.name)

    return docs


def load_docx(path: Path) -> List[Document]:
    from docx import Document as DocxDocument

    docx = DocxDocument(str(path))
    full = "\n\n".join(p.text for p in docx.paragraphs if p.text.strip())
    full = clean_text(full)
    return [Document(page_content=full, metadata={"page": 1})] if full else []


def load_txt(path: Path) -> List[Document]:
    text = clean_text(path.read_text(encoding="utf-8", errors="ignore"))
    return [Document(page_content=text, metadata={"page": 1})] if text else []


def load_image(path: Path) -> List[Document]:
    """Load an image file (PNG/JPG/TIFF/…) via Tesseract OCR."""
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        logger.error(
            "pytesseract or Pillow not installed — cannot OCR image %s. "
            "Run: pip install pytesseract Pillow  and install Tesseract on your system.",
            path.name,
        )
        return []

    try:
        img = Image.open(str(path))
        text = clean_text(pytesseract.image_to_string(img, lang="vie+eng"))
        if not text.strip():
            logger.warning("OCR returned empty text for image: %s", path.name)
            return []
        logger.info("OCR extracted %d chars from image %s", len(text), path.name)
        return [Document(page_content=text, metadata={"page": 1})]
    except Exception as e:
        logger.exception("Failed to OCR image %s: %s", path.name, e)
        return []


def load_document(path: Path) -> List[Document]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return load_pdf(path)
    elif suffix == ".docx":
        return load_docx(path)
    elif suffix == ".txt":
        return load_txt(path)
    elif suffix in _IMAGE_EXTS:
        return load_image(path)
    raise ValueError(f"Unsupported file type: {suffix}")


# ── Chunking ───────────────────────────────────────────────────────────────

def chunk_documents(docs: List[Document], metadata: dict) -> List[Document]:
    splitter = RecursiveCharacterTextSplitter(
        separators=VI_SEPARATORS,
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
        length_function=len,
        is_separator_regex=False,
    )

    chunks = []
    for doc in docs:
        for i, split in enumerate(splitter.split_text(doc.page_content)):
            chunks.append(Document(
                page_content=split,
                metadata={**metadata, "page": doc.metadata.get("page", 1), "chunk_index": i},
            ))
    return chunks


# ── Pipeline ──────────────────────────────────────────────────────────────

def ingest_file(file_path: Path, overwrite: bool = False) -> dict:
    filename = file_path.name

    already_exists = document_exists(filename)
    if already_exists and not overwrite:
        return {"status": "skipped", "reason": "already_indexed", "source": filename}

    # Delete old data first when overwriting to avoid duplicate chunks
    if already_exists and overwrite:
        logger.info("Overwrite requested — removing old chunks for %s", filename)
        delete_document_by_source(filename)
        remove_document_from_index(filename)

    logger.info("Ingesting %s", filename)

    try:
        docs = load_document(file_path)
    except Exception as e:
        logger.exception("Failed to load %s", filename)
        return {"status": "error", "reason": str(e), "source": filename}

    if not docs:
        return {"status": "error", "reason": "empty_document", "source": filename}

    file_meta = extract_metadata_from_filename(filename)
    chunks = chunk_documents(docs, file_meta)

    if not chunks:
        return {"status": "error", "reason": "no_chunks_produced", "source": filename}

    texts = [c.page_content for c in chunks]
    embeddings = embed_texts(texts)

    # Upsert vector store then BM25; if vector upsert fails we skip BM25
    try:
        upsert_chunks(chunks, embeddings)
    except Exception as e:
        logger.exception("Qdrant upsert failed for %s", filename)
        return {"status": "error", "reason": f"vector store error: {e}", "source": filename}

    add_chunks_to_index(chunks)

    logger.info("Ingested %s: %d pages, %d chunks", filename, len(docs), len(chunks))
    return {
        "status": "success",
        "source": filename,
        "pages": len(docs),
        "chunks": len(chunks),
    }


def ingest_bytes(filename: str, data: bytes, overwrite: bool = False) -> dict:
    """Save raw bytes to data/ then run ingestion pipeline."""
    settings.data_dir.mkdir(parents=True, exist_ok=True)
    dest = settings.data_dir / filename
    dest.write_bytes(data)
    return ingest_file(dest, overwrite=overwrite)


def ingest_all(data_dir: Optional[Path] = None, overwrite: bool = False) -> List[dict]:
    """Recursively scan data_dir (all subdirectories) and ingest supported files."""
    data_dir = data_dir or settings.data_dir
    files = sorted(
        f for f in data_dir.rglob("*")
        if f.is_file() and f.suffix.lower() in SUPPORTED_EXTS
    )
    if not files:
        logger.info("ingest_all: no supported files found under %s", data_dir)
        return []
    logger.info("ingest_all: found %d file(s) under %s", len(files), data_dir)
    return [ingest_file(f, overwrite=overwrite) for f in files]
